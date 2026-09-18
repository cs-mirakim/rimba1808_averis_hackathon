"""
pipeline/parsers.py — Document parsers for .txt, .docx, .xlsx, and .pdf attachments.

Responsibilities:
1. parse_txt(file_path): Plain text parsing with encoding safety.
2. parse_docx(file_path): Extract paragraphs and tables using python-docx.
3. parse_xlsx(file_path): Extract cells row-by-row across sheets using openpyxl.
4. parse_pdf(file_path): Extract text page-by-page using pypdf.
5. Safety check:
   - If file size is 0 bytes or cannot be parsed/corrupt:
     Raises UnreadableFileError (tagged with status: "NEEDS_REVIEW", review_reason: "unreadable").
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Optional, Union

# Suppress benign pypdf stream pointer warnings
logging.getLogger("pypdf").setLevel(logging.ERROR)


class UnreadableFileError(Exception):
    """Exception raised when a document cannot be read, is empty (0 bytes), or is corrupt."""

    def __init__(self, message: str, file_path: Union[str, Path] = ""):
        super().__init__(message)
        self.file_path = str(file_path)
        self.status = "NEEDS_REVIEW"
        self.review_reason = "unreadable"

    def to_dict(self) -> dict:
        """Return standardized review status dict for downstream modules."""
        return {
            "status": self.status,
            "review_reason": self.review_reason,
            "has_defect": False,
            "defect_fields": [],
            "error_detail": str(self),
            "file_path": self.file_path,
        }


def _verify_file_non_empty(file_path: Union[str, Path]) -> Path:
    """Verify that file exists and is strictly greater than 0 bytes."""
    path = Path(file_path)
    if not path.exists():
        raise UnreadableFileError(f"File not found: {path}", file_path=path)

    try:
        size = path.stat().st_size
    except OSError as e:
        raise UnreadableFileError(f"Unable to access file stats: {e}", file_path=path)

    if size == 0:
        raise UnreadableFileError(f"File is empty (0 bytes): {path}", file_path=path)

    return path


def parse_txt(file_path: Union[str, Path]) -> str:
    """
    Parse a plain text file (.txt).

    Args:
        file_path: Path to the .txt file.

    Returns:
        Clean extracted text string.

    Raises:
        UnreadableFileError: If file is 0 bytes or cannot be decoded/opened.
    """
    path = _verify_file_non_empty(file_path)

    try:
        # Try UTF-8 first with replace fallback to prevent unexpected crashes on exotic bytes
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
            if not content.strip():
                raise UnreadableFileError(f"File contains no readable text: {path}", file_path=path)
            return content.strip()
    except UnreadableFileError:
        raise
    except Exception as e:
        raise UnreadableFileError(f"Failed to read text file {path.name}: {e}", file_path=path) from e


def parse_docx(file_path: Union[str, Path]) -> str:
    """
    Parse a Microsoft Word document (.docx) extracting paragraphs and tables.

    Args:
        file_path: Path to the .docx file.

    Returns:
        Extracted text including table cell representations.

    Raises:
        UnreadableFileError: If file is 0 bytes or corrupt/unopenable.
    """
    path = _verify_file_non_empty(file_path)

    try:
        import docx
    except ImportError as e:
        raise RuntimeError("python-docx library is required: pip install python-docx") from e

    try:
        doc = docx.Document(path)
        chunks = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                chunks.append(text)

        # Extract text from tables (crucial for shipping instructions and BLs)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    chunks.append(" | ".join(row_cells))

        extracted_text = "\n".join(chunks).strip()
        if not extracted_text:
            raise UnreadableFileError(f"DOCX document contains no readable text: {path}", file_path=path)

        return extracted_text

    except UnreadableFileError:
        raise
    except Exception as e:
        raise UnreadableFileError(f"Corrupt or invalid DOCX file {path.name}: {e}", file_path=path) from e


def parse_xlsx(file_path: Union[str, Path]) -> str:
    """
    Parse an Excel spreadsheet (.xlsx) extracting cell values across all sheets row-by-row.

    Args:
        file_path: Path to the .xlsx file.

    Returns:
        Tabular formatted string representation of spreadsheet contents.

    Raises:
        UnreadableFileError: If file is 0 bytes or corrupt/unopenable.
    """
    path = _verify_file_non_empty(file_path)

    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("openpyxl library is required: pip install openpyxl") from e

    try:
        # data_only=True ensures calculated formula values are read instead of raw formulas
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        lines = []

        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            lines.append(f"--- SHEET: {sheetname} ---")
            for row in ws.iter_rows(values_only=True):
                # Format non-empty cells
                row_vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if row_vals:
                    lines.append(" | ".join(row_vals))

        wb.close()

        extracted_text = "\n".join(lines).strip()
        if not extracted_text:
            raise UnreadableFileError(f"XLSX spreadsheet contains no readable text: {path}", file_path=path)

        return extracted_text

    except UnreadableFileError:
        raise
    except Exception as e:
        raise UnreadableFileError(f"Corrupt or invalid XLSX file {path.name}: {e}", file_path=path) from e


def parse_pdf(file_path: Union[str, Path]) -> str:
    """
    Parse a Portable Document Format (.pdf) extracting text page-by-page.

    Args:
        file_path: Path to the .pdf file.

    Returns:
        Extracted text from all pages.

    Raises:
        UnreadableFileError: If file is 0 bytes or corrupt/unopenable.
    """
    path = _verify_file_non_empty(file_path)

    try:
        import pypdf
    except ImportError as e:
        raise RuntimeError("pypdf library is required: pip install pypdf") from e

    try:
        reader = pypdf.PdfReader(path)
        if len(reader.pages) == 0:
            raise UnreadableFileError(f"PDF contains 0 pages: {path}", file_path=path)

        pages_text = []
        for idx, page in enumerate(reader.pages, start=1):
            page_content = page.extract_text() or ""
            if page_content.strip():
                pages_text.append(page_content.strip())

        extracted_text = "\n\n".join(pages_text).strip()
        if not extracted_text:
            raise UnreadableFileError(f"PDF contains no extractable text: {path}", file_path=path)

        return extracted_text

    except UnreadableFileError:
        raise
    except Exception as e:
        raise UnreadableFileError(f"Corrupt or unreadable PDF file {path.name}: {e}", file_path=path) from e
